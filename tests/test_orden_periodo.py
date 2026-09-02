"""Tests del orden A-Z por COD_PERIODO dentro de la rutina de limpieza.

Fijan lo que hace util el ordenado: que cada periodo ocupe UN bloque contiguo
(para no cambiar el filtro de Periodo de ISEF07 linea por linea), que los
colores de la Fase 1 viajen con su fila, y que los numeros de fila del
`ResultadoValidacion` sigan apuntando al archivo que se acaba de escribir.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from moodle_sinu.constantes import (  # noqa: E402
    COLUMNAS_ESPERADAS,
    COL_COD_PERIODO,
    COL_VALIDACION_RPA,
    Color,
    Validacion,
)
from moodle_sinu.modelos import FilaReporte  # noqa: E402
from moodle_sinu.orden_periodo import bloques_periodo, clave_periodo  # noqa: E402
from moodle_sinu.plan_vinculacion import construir_plan  # noqa: E402
from moodle_sinu.salida_excel import escribir_copia_coloreada  # noqa: E402
from moodle_sinu.validacion import validar_reporte  # noqa: E402

#: Columna donde cae VALIDACION_RPA: la que sigue a las 15 del reporte.
LETRA_RPA = "P"

# Misma fila base que test_validacion, para no divergir.
FILA_OK = [
    "1000000101",
    "BOHORQUEZ ROMERO ALEJANDRA",
    "ALEJANDRA.BOHORQUEZR@CUN.EDU.CO",
    "3142939689",
    "3142939689",
    "IV001",
    "IN001",
    "A1I01",
    "806589",
    "50609",
    "B4",
    "PAGO",
    "26I04",
    "A1I01/50609/26I04/B4/IV001/IN001",
    "NO_MATRICULADO",
]

_INDICE = {campo: i for i, campo in enumerate(COLUMNAS_ESPERADAS)}


def fila(**cambios) -> list:
    """Copia de FILA_OK con campos sobreescritos por nombre de columna."""
    valores = list(FILA_OK)
    for campo, valor in cambios.items():
        valores[_INDICE[campo]] = valor
    return valores


def escribir_xlsx(tmp_path: Path, filas: list[list], *, pie: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(list(COLUMNAS_ESPERADAS))
    for f in filas:
        ws.append(f)
    if pie:
        ws.append(["Filtros aplicados: pago es PAGO"] + [None] * 14)
    ruta = tmp_path / "reporte.xlsx"
    wb.save(ruta)
    wb.close()
    return ruta


def _columna(ws, letra: str, desde: int, hasta: int) -> list:
    return [ws[f"{letra}{n}"].value for n in range(desde, hasta + 1)]


def test_ordena_los_periodos_de_la_a_a_la_z(tmp_path):
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="2", COD_PERIODO="26I04"),
            fila(IDENTIFICACION="3", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="4", COD_PERIODO="26A01"),
            fila(IDENTIFICACION="5", COD_PERIODO="26I04"),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        ws = wb["Export"]
        assert _columna(ws, "M", 2, 6) == ["26A01", "26I04", "26I04", "26V05", "26V05"]
    finally:
        wb.close()


def test_cada_periodo_queda_en_un_solo_bloque_contiguo(tmp_path):
    """Es el objetivo del cambio: un bloque por periodo, no uno por linea."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION=str(n), COD_PERIODO=p)
            for n, p in enumerate(["26V05", "26I04", "26V05", "26I04", "26V05"], 1)
        ],
    )
    resultado = validar_reporte(entrada)
    assert len(bloques_periodo(resultado)) == 5  # antes: un bloque por fila

    escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")
    assert bloques_periodo(resultado) == [("26I04", 2), ("26V05", 3)]


def test_dentro_de_un_periodo_se_conserva_el_orden_de_aparicion(tmp_path):
    """El orden estable es lo que permite reanudar una corrida interrumpida."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="A", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="B", COD_PERIODO="26I04"),
            fila(IDENTIFICACION="C", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="D", COD_PERIODO="26V05"),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        assert _columna(wb["Export"], "A", 2, 5) == ["B", "A", "C", "D"]
    finally:
        wb.close()


def test_los_colores_viajan_con_su_fila(tmp_path):
    """Una celda naranja tiene que seguir siendo naranja tras moverse."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            # Sin CORREO -> naranja en la columna C y fila omitida.
            fila(IDENTIFICACION="2", COD_PERIODO="26I04", CORREO=""),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        ws = wb["Export"]
        # La fila sin correo es la del 26I04, que ahora va primera.
        assert ws["A2"].value == "2"
        assert ws["C2"].fill.start_color.rgb.endswith(Color.NARANJA.value)
        assert ws[f"{LETRA_RPA}2"].value == Validacion.DATO_INCOMPLETO.value
        # La fila que no tenia hallazgos sigue sin pintar.
        assert ws["A3"].value == "1"
        assert ws["C3"].fill.fill_type in (None, "none")
    finally:
        wb.close()


def test_los_numeros_de_fila_apuntan_al_archivo_escrito(tmp_path):
    """Tras ordenar, `fila.fila` describe la copia, no el original."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="2", COD_PERIODO="26I04"),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        ws = wb["Export"]
        for f in resultado.filas:
            assert ws[f"A{f.fila}"].value == f.identificacion
    finally:
        wb.close()


def test_los_hallazgos_se_renumeran_igual_que_sus_filas(tmp_path):
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(
                IDENTIFICACION="2",
                COD_PERIODO="26I04",
                NOMBRE_COMPLETO="  JUAN  PEREZ ",
            ),
        ],
    )
    resultado = validar_reporte(entrada)
    escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    azules = resultado.hallazgos_por_color(Color.AZUL_CLARO)
    assert azules, "la fila con espacios tenia que marcarse en azul"
    por_fila = {f.fila: f for f in resultado.filas}
    for h in azules:
        assert por_fila[h.fila].identificacion == "2"


def test_el_pie_de_power_bi_no_se_mueve(tmp_path):
    """La fila 'Filtros aplicados:' no es un registro: se queda donde estaba."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="2", COD_PERIODO="26I04"),
        ],
        pie=True,
    )
    resultado = validar_reporte(entrada)
    assert resultado.filas_pie_descartadas == [4]
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        ws = wb["Export"]
        assert _columna(ws, "M", 2, 3) == ["26I04", "26V05"]
        assert str(ws["A4"].value).startswith("Filtros aplicados:")
    finally:
        wb.close()


def test_las_filas_sin_periodo_van_al_final(tmp_path):
    """No son ejecutables (no hay filtro de Periodo que fijar): estorban arriba."""
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO=""),
            fila(IDENTIFICACION="2", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="3", COD_PERIODO="26A01"),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    wb = load_workbook(destino)
    try:
        assert _columna(wb["Export"], "A", 2, 4) == ["3", "2", "1"]
    finally:
        wb.close()


def test_sin_ordenar_deja_el_orden_original(tmp_path):
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="2", COD_PERIODO="26A01"),
        ],
    )
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(
        resultado, tmp_path / "salida.xlsx", ordenar_periodo=False
    )

    wb = load_workbook(destino)
    try:
        assert _columna(wb["Export"], "A", 2, 3) == ["1", "2"]
    finally:
        wb.close()


def test_el_orden_del_archivo_coincide_con_el_del_plan(tmp_path):
    """El .xlsx subido y el plan tienen que recorrer los periodos igual.

    Es la razon del cambio: si el operador mira el Sheet en un orden y el robot
    trabaja en otro, no puede seguir la corrida.
    """
    entrada = escribir_xlsx(
        tmp_path,
        [
            fila(IDENTIFICACION="1", COD_PERIODO="26V05"),
            fila(IDENTIFICACION="2", COD_PERIODO="26A01"),
            fila(IDENTIFICACION="3", COD_PERIODO="26I04"),
        ],
    )
    resultado = validar_reporte(entrada)
    escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")

    plan = construir_plan(resultado)
    assert [lote.cod_periodo for lote in plan.lotes] == [
        periodo for periodo, _ in bloques_periodo(resultado)
    ]


def test_clave_periodo_normaliza_espacios_y_mayusculas():
    def _f(periodo):
        return FilaReporte(fila=2, crudos={}, saneados={COL_COD_PERIODO: periodo})

    assert clave_periodo(_f("26i04")) == clave_periodo(_f(" 26I04 "))
    assert clave_periodo(_f("")) > clave_periodo(_f("ZZZZZ"))


def test_revalidar_la_copia_no_convierte_duplicados_exactos_en_amarillos(tmp_path):
    """Regresion: la columna VALIDACION_RPA no debe entrar en la huella.

    Su texto difiere entre las apariciones de un duplicado exacto ('' en la
    primera, 'OMITIDA: ...' en las demas), asi que al revalidar la copia la
    regla 3 las tomaba por 'datos distintos' y omitia dos filas procesables.
    """
    assert COL_VALIDACION_RPA not in COLUMNAS_ESPERADAS

    repetida = fila(IDENTIFICACION="9", COD_PERIODO="26I04")
    entrada = escribir_xlsx(tmp_path, [repetida, list(repetida)])
    resultado = validar_reporte(entrada)
    destino = escribir_copia_coloreada(resultado, tmp_path / "salida.xlsx")
    assert len(resultado.procesables) == 1  # se deduplica, no se marca amarillo

    segunda = validar_reporte(destino)
    assert not segunda.hallazgos_por_color(Color.AMARILLO)
    assert len(segunda.procesables) == 1
