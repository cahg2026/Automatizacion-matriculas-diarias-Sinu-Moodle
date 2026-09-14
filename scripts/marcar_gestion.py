"""Pinta en el .xlsx la gestion de la etapa 4: verde vinculado, rojo lo demas.

Regla del dueno del proceso (04/09/2026): la gestion debe quedar reflejada por
cada unidad de trabajo **(cedula, materia)** procesada.

    VERDE  -> vinculada y con el check de Moodle confirmado
    ROJO   -> fallo, sin correspondencia en SINU, o check sin confirmar

Por que sobre el .xlsx y no sobre el Sheet en vivo
--------------------------------------------------
No hay API de Google en este proyecto: la organizacion no concede acceso a
Google Cloud Console, asi que no existe cuenta de servicio ni `client_secret`
(ver `requirements.txt`). Y pintar celdas por RPA no es viable de forma fiable:
Google Sheets dibuja la cuadricula en un `<canvas>` y las celdas **no existen en
el DOM** -- por eso el periodo del dia se lee exportando el documento a CSV en
vez de raspar la tabla (ver `periodo_sheet`).

Asi que se pinta el .xlsx, que es el mismo archivo que se subio como Sheet, y se
vuelve a subir. No es tiempo real: es una foto. La alternativa en vivo requiere
la API.

De donde salen los veredictos
-----------------------------
De `logs/resultados_etapa4.jsonl` (ver `diario_resultados`), que el CLI de la
etapa 4 escribe tras CADA operacion. No se parsean logs.

La guarda de la fila
--------------------
El diario guarda el numero de fila, y el Sheet conserva el orden del .xlsx
porque se sube ya ordenado A-Z. Pero antes de pintar se comprueba que la fila
contenga de verdad esa cedula y esa materia. Pintar la fila equivocada es del
mismo tipo de error que ejecutar sobre la matricula equivocada, asi que ante una
discrepancia NO se pinta: se reporta.

Uso:
    python scripts/marcar_gestion.py <reporte_validado.xlsx> [--salida X.xlsx]
                                     [--subir] [--visible]
"""

from __future__ import annotations

import argparse
import logging
import re
import shutil
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

from moodle_sinu import diario_resultados as diario  # noqa: E402
from moodle_sinu import registro  # noqa: E402
from moodle_sinu.config import Config, asegurar_directorios  # noqa: E402
from moodle_sinu.constantes import (  # noqa: E402
    COL_COD_MATERIA,
    COL_COD_PERIODO,
    COL_IDENTIFICACION,
    COL_NUM_GRUPO,
    COL_VALIDACION_RPA,
    Color,
)

log = logging.getLogger("marcar_gestion")

#: Los dos colores que pidio el dueno del proceso. Ya estaban en la paleta como
#: colores de Fase 2, asi que no se inventa nada nuevo.
VERDE = Color.VERDE_CLARO
ROJO = Color.ROJO_CLARO


#: `reporte_moodle_sinu_20260903_120929...` -> 2026-09-03. Ese sello lo pone la
#: etapa 1 al descargar, asi que es la fecha del REPORTE.
RX_SELLO = re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)")


def fecha_del_reporte(archivo: Path) -> date | None:
    """La fecha del reporte, leida del sello de su nombre de archivo.

    Regla del dueno del proceso (04/09/2026): el nombre en Drive lleva **la
    fecha del reporte**, no la del dia en que se ejecuta la gestion. Son
    distintas siempre que se procese al dia siguiente -- paso el 04/09 con el
    reporte del 03/09, y el archivo salio como `REPORTE 04-09-2026`.
    """
    m = RX_SELLO.search(archivo.name)
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _copia_fuera_de_onedrive(origen: Path) -> Path:
    """Copia el archivo al temporal del sistema y devuelve la ruta.

    `data/processed/` vive dentro de OneDrive, y OneDrive bloquea los archivos
    mientras los sincroniza -- el README ya lo advierte para el perfil del
    navegador y el `.venv`. El 04/09/2026 la subida de un archivo recien escrito
    ahi fallo CON ventana ("no aparecio en la carpeta tras 300s"), el mismo
    sintoma que da un archivo que Chrome no puede leer entero.

    La copia que se entrega al selector de archivos se escribe fuera del alcance
    del sincronizador. Es barato y descarta una causa entera.
    """
    base = Path(tempfile.gettempdir()) / "moodle-sinu-subidas"
    base.mkdir(parents=True, exist_ok=True)
    destino = base / origen.name
    shutil.copy2(origen, destino)
    return destino


def _argumentos() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Marca la gestion de la etapa 4 en el .xlsx")
    p.add_argument("reporte", help="El .xlsx validado (el que se subio como Sheet)")
    p.add_argument(
        "--salida",
        default=None,
        help="Ruta del .xlsx marcado. Por defecto, junto al original con sello "
        "'_gestion_<fecha>'. NUNCA sobreescribe el original.",
    )
    p.add_argument(
        "--subir",
        action="store_true",
        help="Sube el resultado a Drive convertido a Sheet. Sin esto solo escribe "
        "el archivo local, que es lo que conviene la primera vez.",
    )
    p.add_argument(
        "--visible",
        action="store_true",
        help="Navegador con ventana. OBLIGATORIO en la practica para subir: la "
        "subida a Drive falla en headless (comprobado el 03/09/2026).",
    )
    p.add_argument(
        "--fecha",
        default=None,
        metavar="YYYY-MM-DD",
        help="Fecha del reporte para el nombre en Drive. Por defecto se deduce "
        "del sello del archivo, que es lo correcto: el nombre lleva la fecha del "
        "REPORTE, no la del dia en que corre la gestion.",
    )
    p.add_argument(
        "--consecutivo",
        type=int,
        default=None,
        help="Fuerza el numero del reporte, para que el marcado conserve el mismo "
        "que el original.",
    )
    p.add_argument(
        "--reemplazar",
        action="store_true",
        help="Si ya hay un reporte de esa fecha en Drive, lo APARTA (le anade un "
        "sufijo) y el nombre canonico pasa al marcado. No borra nada y no rompe "
        "enlaces: las URL de Drive van por id, no por nombre.",
    )
    p.add_argument(
        "--traza",
        action="store_true",
        help="Guarda una traza de Playwright en logs/. Es lo que hay que mirar "
        "cuando la subida dice 'no aparecio en la carpeta': la traza muestra si "
        "el archivo se entrego al selector y que hizo Drive despues.",
    )
    p.add_argument("--verbose", "-v", action="store_true")
    return p.parse_args()


def _indices(ws) -> dict[str, int]:
    """Columna (1-based) de cada encabezado que hace falta."""
    encabezados = {}
    for col, celda in enumerate(ws[1], start=1):
        if celda.value:
            encabezados[str(celda.value).strip()] = col
    faltan = [
        c
        for c in (COL_IDENTIFICACION, COL_COD_MATERIA, COL_COD_PERIODO)
        if c not in encabezados
    ]
    if faltan:
        raise SystemExit(
            f"El .xlsx no trae las columnas {faltan}. Encabezados vistos: "
            f"{sorted(encabezados)}"
        )
    return encabezados


def _texto(ws, fila: int, col: int | None) -> str:
    if not col:
        return ""
    return str(ws.cell(row=fila, column=col).value or "").strip()


def main() -> int:
    args = _argumentos()
    asegurar_directorios()
    registro.configurar(
        logging.DEBUG if args.verbose else logging.INFO, etiqueta="marcar_gestion"
    )

    entrada = Path(args.reporte)
    if not entrada.is_file():
        raise SystemExit(f"No existe {entrada}")

    unidades = diario.ultimo_por_unidad()
    if not unidades:
        print(f"El diario esta vacio ({diario.RUTA_RESULTADOS}). Nada que marcar.")
        return 0

    wb = load_workbook(entrada)
    ws = wb.active
    cols = _indices(ws)
    col_rpa = cols.get(COL_VALIDACION_RPA)
    if not col_rpa:
        col_rpa = ws.max_column + 1
        ws.cell(row=1, column=col_rpa, value=COL_VALIDACION_RPA)
    ultima_col = ws.max_column

    verdes = rojos = 0
    descartadas: list[str] = []

    for (periodo, cedula, objetivo), apunte in unidades.items():
        fila = apunte.get("fila") or 0
        if not fila:
            descartadas.append(f"{cedula} / {objetivo}: el apunte no trae fila")
            continue

        # --- la guarda: que la fila sea de verdad esta unidad ---
        ced_fila = _texto(ws, fila, cols.get(COL_IDENTIFICACION))
        mat_fila = _texto(ws, fila, cols.get(COL_COD_MATERIA))
        gru_fila = _texto(ws, fila, cols.get(COL_NUM_GRUPO))
        esperado = apunte.get("cod_materia", "")
        if ced_fila != cedula or mat_fila.upper() != esperado.upper():
            descartadas.append(
                f"{cedula} / {objetivo}: la fila {fila} contiene "
                f"{ced_fila} / {mat_fila}/{gru_fila}. NO se pinta."
            )
            continue

        es_verde = apunte.get("veredicto") == diario.Veredicto.VERDE.value
        relleno = PatternFill(
            start_color=(VERDE if es_verde else ROJO).value,
            end_color=(VERDE if es_verde else ROJO).value,
            fill_type="solid",
        )
        for col in range(1, ultima_col + 1):
            ws.cell(row=fila, column=col).fill = relleno

        etiqueta = "VINCULADA Y VALIDADA" if es_verde else "REVISAR"
        ws.cell(
            row=fila,
            column=col_rpa,
            value=f"{etiqueta} - {apunte.get('motivo', '')} ({apunte.get('momento', '')[:16]})",
        )
        if es_verde:
            verdes += 1
        else:
            rojos += 1

    salida = (
        Path(args.salida)
        if args.salida
        else entrada.with_name(
            f"{entrada.stem}_gestion_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
    )
    wb.save(salida)

    print()
    print(f"Unidades en el diario : {len(unidades)}")
    print(f"  pintadas VERDE      : {verdes}")
    print(f"  pintadas ROJO       : {rojos}")
    if descartadas:
        print(f"  NO pintadas         : {len(descartadas)}")
        for d in descartadas[:10]:
            print(f"     {d}")
    print(f"Archivo marcado       : {salida}")

    if not args.subir:
        print()
        print("No se subio (falta --subir). Revisar el archivo y despues:")
        print(f"  python scripts/marcar_gestion.py {entrada} --subir --visible")
        return 0

    # La subida reutiliza la maquinaria probada de la etapa 2.
    from moodle_sinu.subidor_drive import ErrorSubidaDrive, subir_reporte

    cfg = Config.desde_entorno()

    if args.fecha:
        fecha = date.fromisoformat(args.fecha)
    else:
        fecha = fecha_del_reporte(entrada)
        if fecha is None:
            print()
            print(
                f"!! No se pudo deducir la fecha del reporte de '{entrada.name}'. "
                "Indicarla con --fecha YYYY-MM-DD: el nombre en Drive debe llevar "
                "la fecha del reporte, no la de hoy."
            )
            return 2
    print(f"Fecha para Drive      : {fecha:%d/%m/%Y} (del reporte, no de hoy)")

    a_subir = _copia_fuera_de_onedrive(salida)
    log.info("Copia para subir, fuera de OneDrive: %s", a_subir)

    try:
        r = subir_reporte(
            a_subir,
            cfg=cfg,
            fecha=fecha,
            consecutivo=args.consecutivo,
            headless=False if args.visible else None,
            abrir_al_terminar=False,
            reemplazar=args.reemplazar,
            con_traza=args.traza,
        )
    except ErrorSubidaDrive as exc:
        log.error("No se pudo subir: %s", exc)
        print(f"\n!! La subida fallo: {exc}")
        print("   El archivo local SI esta escrito; se puede subir a mano.")
        return 1

    print()
    print(f"Subido como : {getattr(r, 'nombre', '(ver log)')}")
    print(f"Sheet       : {getattr(r, 'url_sheet', '(ver log)')}")
    for aviso in getattr(r, "advertencias", []):
        print(f"  aviso: {aviso}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
