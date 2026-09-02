"""Etapa 1: exportacion automatica del reporte desde Power BI (Playwright).

No hay API para "Datos con diseno actual" -- la exportacion de un visual con el
diseno y los filtros aplicados solo existe en la interfaz -- asi que esta etapa
es RPA sobre el navegador. Los selectores viven en `selectores_powerbi.py`,
extraidos de una grabacion real; aqui solo esta la orquestacion, los reintentos
y el manejo de fallos.

La exportacion es de solo lectura: no modifica nada en Power BI, por lo que
`MODO_SIMULACION` no la afecta (esa bandera protege la accion irreversible
"Vincular grupos matriculados" de la etapa 4).
"""

from __future__ import annotations

import logging
import re
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from urllib.parse import urlsplit

from openpyxl import load_workbook
from playwright.sync_api import (
    BrowserContext,
    Download,
    Error as ErrorPlaywright,
    Locator,
    Page,
    TimeoutError as ErrorTiempoPlaywright,
    sync_playwright,
)

from . import selectores_powerbi as sel
from .config import DIR_CRUDO, DIR_LOGS, Config
from .constantes import COLUMNAS_ESPERADAS
from .navegador import abrir_contexto

log = logging.getLogger(__name__)


class ErrorExportacionPowerBI(RuntimeError):
    """La exportacion no se pudo completar.

    Se lanza en vez de continuar con un archivo parcial o inexistente: el resto
    del proceso asume que el .xlsx descargado es el reporte del dia.
    """


@dataclass
class ResultadoExportacion:
    """Salida de la etapa 1."""

    archivo: Path
    """Ruta del .xlsx descargado."""

    bytes: int
    segundos: float
    nombre_sugerido: str
    """Nombre con el que Power BI ofrecio la descarga."""

    ruta_navegacion: str
    """Como se llego al informe: 'url-directa', 'busqueda' o 'area-de-trabajo'."""

    advertencias: list[str] = field(default_factory=list)
    """Pasos que no se pudieron confirmar. No abortan, pero hay que revisarlos."""

    traza: Path | None = None
    """Traza de Playwright, si se pidio (`playwright show-trace <ruta>`)."""

    no_matriculado: int | None = None
    """Valor de la tarjeta 'NO MATRICULADO' del tablero. Es el numero que va en
    el nombre del reporte en Drive: `REPORTE <fecha> #<no_matriculado>`. NO es
    un consecutivo (decision del dueno del proceso, 24/08/2026)."""

    filas_datos: int | None = None
    """Filas de datos del .xlsx descargado. Debe coincidir con `no_matriculado`:
    la vista de Power BI viene filtrada a VALIDACION != MATRICULADO, asi que la
    tarjeta y la tabla cuentan lo mismo. Si difieren, algo fallo."""

    captura_dialogo: Path | None = None
    """Captura del dialogo de exportacion antes de confirmar, si se pidio.
    Es la evidencia de que 'Datos con diseno actual' quedo seleccionado."""


# ---------------------------------------------------------------------------
# Utilidades de localizacion
# ---------------------------------------------------------------------------


def _ms(segundos: float) -> float:
    return segundos * 1000


def _visible(locator: Locator, timeout_seg: float) -> bool:
    """True si el locator llega a ser visible dentro del plazo.

    Envuelve el timeout porque en este flujo la ausencia de un elemento es a
    menudo informacion valida (el dialogo de sesion no aparece si el perfil ya
    la recuerda), no un fallo. Tambien absorbe los errores de navegacion: el
    SPA de Power BI redirige por su cuenta y puede destruir el contexto de
    ejecucion mientras esperamos, lo que tampoco significa "no esta".
    """
    try:
        locator.first.wait_for(state="visible", timeout=_ms(timeout_seg))
        return True
    except ErrorTiempoPlaywright:
        return False
    except ErrorPlaywright as exc:
        log.debug("Espera interrumpida por la navegacion del SPA: %s", exc)
        return False


def _es_aborto_de_navegacion(exc: BaseException) -> bool:
    """True si el error es un goto abortado por otra navegacion en curso.

    Power BI arranca navegaciones de cliente por su cuenta -- sobre todo justo
    despues del login -- y cuando una de ellas pisa nuestro `goto`, Chromium lo
    cancela con net::ERR_ABORTED. Nuestra navegacion fallo, pero la pagina
    normalmente acaba donde queriamos ir, asi que no es motivo para abortar.
    """
    return "ERR_ABORTED" in str(exc)


def _mismo_origen(actual: str, esperado: str) -> bool:
    a, e = urlsplit(actual), urlsplit(esperado)
    return (a.scheme, a.netloc) == (e.scheme, e.netloc)


def _en_destino(actual: str, esperado: str, exigir_ruta: bool) -> bool:
    """Comprueba si la pagina acabo en el destino pretendido.

    Con `exigir_ruta=False` basta el mismo origen: cualquier pagina del shell de
    Power BI sirve, porque la barra de busqueda global esta en todas. Para la
    URL de un informe concreto la ruta si importa -- dar por bueno otro informe
    llevaria a exportar datos equivocados.
    """
    if not _mismo_origen(actual, esperado):
        return False
    if not exigir_ruta:
        return True

    ruta_actual = urlsplit(actual).path.rstrip("/")
    ruta_esperada = urlsplit(esperado).path.rstrip("/")
    if ruta_actual == ruta_esperada:
        return True

    # Power BI reescribe el ultimo segmento de la ruta: la seccion/pagina del
    # informe cambia sola al cargar (…/reports/<id>/ReportSection ->
    # …/reports/<id>/ReportSection2). Lo que no puede cambiar es el informe.
    padre = ruta_esperada.rsplit("/", 1)[0]
    if not padre:
        return False
    return ruta_actual == padre or ruta_actual.startswith(padre + "/")


def _navegar(
    page: Page,
    url: str,
    cfg: Config,
    *,
    descripcion: str,
    exigir_ruta: bool = True,
) -> None:
    """`page.goto` tolerante a las navegaciones propias del SPA.

    Raises:
        ErrorExportacionPowerBI: si la navegacion falla por cualquier motivo que
            no sea un aborto, o si tras los reintentos la pagina no esta en el
            destino.
    """
    for intento in (1, 2):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=_ms(cfg.timeout_render_seg))
            return
        except ErrorPlaywright as exc:
            if not _es_aborto_de_navegacion(exc):
                raise ErrorExportacionPowerBI(
                    f"No se pudo abrir {descripcion} ({url}): {exc}"
                ) from exc

            # Dejar que Power BI termine la navegacion que inicio el mismo.
            try:
                page.wait_for_load_state(
                    "domcontentloaded", timeout=_ms(cfg.timeout_operacion_seg)
                )
            except ErrorPlaywright:
                pass

            if _en_destino(page.url, url, exigir_ruta):
                log.info(
                    "Power BI interrumpio la navegacion a %s con su propia "
                    "redireccion, pero la pagina ya estaba en el destino (%s).",
                    descripcion,
                    page.url,
                )
                return

            log.warning(
                "La navegacion a %s la interrumpio el propio Power BI "
                "(intento %d/2); la pagina quedo en %s. Reintentando.",
                descripcion,
                intento,
                page.url,
            )

    raise ErrorExportacionPowerBI(
        f"Power BI aborto dos veces la navegacion a {descripcion} ({url}); "
        f"la pagina quedo en {page.url}."
    )


def _primero_visible(
    candidatos: list[tuple[str, Locator]],
    timeout_seg: float,
    descripcion: str,
) -> tuple[str, Locator]:
    """Devuelve el primer candidato visible, en orden de preferencia.

    Los candidatos van del selector grabado al de reserva por rol accesible.
    Reparte el plazo para que probar una reserva no multiplique la espera.
    """
    if not candidatos:  # pragma: no cover - error de programacion
        raise ValueError("sin candidatos")

    por_intento = max(timeout_seg / len(candidatos), 2.0)
    for etiqueta, locator in candidatos:
        if _visible(locator, por_intento):
            if etiqueta != candidatos[0][0]:
                log.warning(
                    "%s: el selector grabado no aparecio, se uso la reserva '%s'. "
                    "Puede haber cambiado el idioma o la version de la UI.",
                    descripcion,
                    etiqueta,
                )
            return etiqueta, locator

    raise ErrorExportacionPowerBI(
        f"No se encontro {descripcion}. Selectores probados: "
        + ", ".join(e for e, _ in candidatos)
        + ". Volver a grabar el flujo con 'playwright codegen https://app.powerbi.com/'."
    )


# ---------------------------------------------------------------------------
# Paso 1: inicio de sesion
# ---------------------------------------------------------------------------


def verificar_credenciales(cfg: Config) -> None:
    """Falla si no hay credenciales de Power BI.

    Se comprueba antes de abrir el navegador: arrancar Chromium para descubrir
    que falta una variable de entorno es puro desperdicio.
    """
    faltantes = [
        nombre
        for nombre, valor in (
            ("POWERBI_USUARIO", cfg.powerbi_usuario),
            ("POWERBI_PASSWORD", cfg.powerbi_password),
        )
        if not valor
    ]
    if faltantes:
        raise ErrorExportacionPowerBI(
            f"Faltan {' y '.join(faltantes)} en config/.env "
            "(plantilla en config/.env.example)."
        )


def _iniciar_sesion(page: Page, cfg: Config, advertencias: list[str]) -> None:
    """Autentica contra Power BI via el SSO de Microsoft.

    Tolera que la sesion ya este abierta (perfil persistente): si el campo de
    correo no aparece, se asume autenticado y se sigue sin pedir credenciales.
    """
    log.info("Abriendo %s", sel.URL_INICIO)
    # exigir_ruta=False: la entrada redirige a /singleSignOn o directamente a
    # /home segun si el perfil ya trae sesion.
    _navegar(page, sel.URL_INICIO, cfg, descripcion="la entrada de Power BI", exigir_ruta=False)

    candidatos = [
        (page.get_by_placeholder(sel.PLACEHOLDER_CORREO), cfg.timeout_operacion_seg),
        (page.get_by_placeholder(sel.RX_PLACEHOLDER_CORREO), 3),
    ]
    campo = next((loc for loc, plazo in candidatos if _visible(loc, plazo)), None)
    if campo is None:
        log.info("No hay formulario de correo: la sesion ya estaba iniciada.")
        return

    # Solo aqui son imprescindibles: con perfil persistente y sesion viva no se
    # llega a este punto.
    verificar_credenciales(cfg)
    assert cfg.powerbi_usuario and cfg.powerbi_password  # garantizado arriba

    log.info("Enviando el usuario %s", cfg.powerbi_usuario)
    campo.fill(cfg.powerbi_usuario)
    # La grabacion continuo con Enter, no con un boton: el formulario de Power BI
    # redirige al SSO al enviar. La navegacion posterior es una redireccion, no
    # un goto reproducible.
    campo.press("Enter")

    _, password = _primero_visible(
        [
            ("placeholder 'Contraseña'", page.get_by_placeholder(sel.PLACEHOLDER_PASSWORD)),
            ("placeholder ~contrase|password", page.get_by_placeholder(sel.RX_PLACEHOLDER_PASSWORD)),
        ],
        cfg.timeout_render_seg,
        "el campo de contrasena del SSO de Microsoft",
    )
    password.fill(cfg.powerbi_password)

    _, boton = _primero_visible(
        [
            ("boton 'Iniciar sesión'", page.get_by_role("button", name=sel.BOTON_INICIAR_SESION)),
            ("boton ~iniciar sesi|sign in", page.get_by_role("button", name=sel.RX_BOTON_INICIAR_SESION)),
        ],
        cfg.timeout_operacion_seg,
        "el boton de inicio de sesion",
    )
    boton.click()

    _resolver_mantener_sesion(page, cfg, advertencias)

    # Si tras el login seguimos en el formulario, las credenciales o un segundo
    # factor bloquearon el acceso. Mejor decirlo que agotar timeouts mas abajo.
    if _visible(page.get_by_placeholder(sel.RX_PLACEHOLDER_PASSWORD), 5):
        raise ErrorExportacionPowerBI(
            "El SSO sigue pidiendo contrasena tras enviarla. Revisar las "
            "credenciales de config/.env o, si la cuenta exige MFA, ejecutar una "
            "vez con --visible y POWERBI_PERFIL_NAVEGADOR configurado para "
            "guardar la sesion."
        )

    log.info("Sesion iniciada.")


def _resolver_mantener_sesion(page: Page, cfg: Config, advertencias: list[str]) -> None:
    """Cierra el dialogo "¿Mantener la sesion iniciada?" si aparece.

    Es opcional: no sale cuando el perfil del navegador ya trae la decision.
    """
    casilla = page.get_by_label(sel.CHECK_NO_VOLVER_A_MOSTRAR)
    if not _visible(casilla, min(cfg.timeout_operacion_seg, 15)):
        casilla = page.get_by_label(sel.RX_CHECK_NO_VOLVER_A_MOSTRAR)
        if not _visible(casilla, 3):
            log.debug("No aparecio el dialogo de mantener sesion.")
            return

    try:
        casilla.first.check()
    except ErrorTiempoPlaywright:
        advertencias.append("No se pudo marcar 'No volver a mostrar' en el dialogo de sesion.")

    try:
        _, si = _primero_visible(
            [
                ("boton 'Sí'", page.get_by_role("button", name=sel.BOTON_SI, exact=True)),
                ("boton ~sí|yes", page.get_by_role("button", name=sel.RX_BOTON_SI)),
            ],
            cfg.timeout_operacion_seg,
            "el boton de confirmacion del dialogo de sesion",
        )
        si.click()
    except ErrorExportacionPowerBI as exc:
        # El dialogo puede autocerrarse; no vale la pena abortar por esto.
        advertencias.append(f"Dialogo de sesion sin confirmar: {exc}")


# ---------------------------------------------------------------------------
# Paso 2: abrir el informe
# ---------------------------------------------------------------------------


def _abrir_informe(context: BrowserContext, page: Page, cfg: Config) -> tuple[Page, str]:
    """Deja abierta la pagina del informe. Devuelve (pagina, ruta_navegacion).

    Se prefiere la URL directa: es el unico camino que no depende de etiquetas
    traducidas ni del orden del listado. Las dos alternativas reproducen la
    grabacion cuando no hay URL configurada.
    """
    if cfg.powerbi_url_informe:
        log.info("Abriendo el informe por URL directa.")
        _navegar(page, cfg.powerbi_url_informe, cfg, descripcion="el informe")
        return page, "url-directa"

    log.warning(
        "POWERBI_URL_INFORME esta vacio: se navegara por la UI, que depende del "
        "idioma de la cuenta. Conviene fijar la URL del informe en config/.env."
    )
    # exigir_ruta=False: tras el login Power BI ya esta redirigiendo por su
    # cuenta y puede dejarnos en /home, /groups/me/list o similar. Cualquiera
    # sirve: la barra de busqueda global vive en el shell, no en una pagina.
    _navegar(page, sel.URL_HOME, cfg, descripcion="la portada", exigir_ruta=False)
    # Dejar constancia de donde aterrizamos: si los selectores del shell fallan,
    # lo primero que hay que saber es que pagina se estaba mirando.
    log.info("Shell de Power BI cargado en %s", page.url)

    try:
        return _abrir_por_busqueda(context, page, cfg), "busqueda"
    except ErrorExportacionPowerBI as exc:
        log.warning("La busqueda global no sirvio (%s); probando el area de trabajo.", exc)
        return _abrir_por_area_trabajo(context, page, cfg), "area-de-trabajo"


def _abrir_por_busqueda(context: BrowserContext, page: Page, cfg: Config) -> Page:
    """Busqueda global -> primer resultado que casa con el nombre del informe."""
    barra = page.get_by_test_id(sel.TESTID_BARRA_BUSQUEDA)
    _, caja = _primero_visible(
        [
            ("barra > tri-search-box", barra.get_by_test_id(sel.TESTID_CAJA_BUSQUEDA)),
            ("searchbox por rol", page.get_by_role("searchbox")),
        ],
        cfg.timeout_operacion_seg,
        "la caja de busqueda global",
    )
    caja.fill(cfg.powerbi_termino_busqueda)
    caja.press("Enter")

    # Preferimos el resultado cuyo nombre casa con el informe. El texto grabado
    # ("del area de trabajo: CUN") casa con cualquier resultado del area y solo
    # se usa como ultimo recurso.
    _, resultado = _primero_visible(
        [
            (f"enlace ~{cfg.powerbi_informe}", page.get_by_role("link", name=cfg.powerbi_informe)),
            (f"texto ~{cfg.powerbi_informe}", page.get_by_text(cfg.powerbi_informe, exact=False)),
            ("texto 'del área de trabajo: CUN'", page.get_by_text(sel.TEXTO_RESULTADO_AREA)),
        ],
        cfg.timeout_operacion_seg,
        f"un resultado de busqueda para '{cfg.powerbi_termino_busqueda}'",
    )
    return _clicar_y_seguir(context, page, resultado, cfg)


def _abrir_por_area_trabajo(context: BrowserContext, page: Page, cfg: Config) -> Page:
    """Barra lateral -> area de trabajo -> fila del informe."""
    _, navbar = _primero_visible(
        [
            ("test-id navbar áreas-de-trabajo", page.get_by_test_id(sel.TESTID_NAVBAR_AREAS_TRABAJO)),
            ("rol ~áreas de trabajo", page.get_by_role("button", name=sel.RX_NAVBAR_AREAS_TRABAJO)),
        ],
        cfg.timeout_operacion_seg,
        "la entrada 'Areas de trabajo' de la barra lateral",
    )
    navbar.click()

    area = page.get_by_role(sel.ROL_AREA_TRABAJO, name=cfg.powerbi_workspace)
    if not _visible(area, cfg.timeout_operacion_seg):
        raise ErrorExportacionPowerBI(
            f"No aparecio el area de trabajo '{cfg.powerbi_workspace}'. "
            "Revisar POWERBI_WORKSPACE y los permisos de la cuenta."
        )
    area.first.click()

    fila = page.get_by_role("row", name=sel.rx_fila_informe(cfg.powerbi_informe))
    _, item = _primero_visible(
        [
            ("fila del listado > item-name", fila.get_by_test_id(sel.TESTID_NOMBRE_ITEM)),
            (f"enlace ~{cfg.powerbi_informe}", page.get_by_role("link", name=cfg.powerbi_informe)),
        ],
        cfg.timeout_render_seg,
        f"el informe '{cfg.powerbi_informe}' en el area '{cfg.powerbi_workspace}'",
    )
    return _clicar_y_seguir(context, page, item, cfg)


def _clicar_y_seguir(context: BrowserContext, page: Page, locator: Locator, cfg: Config) -> Page:
    """Clica y devuelve la pagina donde quedo el informe.

    En la grabacion el informe se abrio en una pestana nueva (`expect_popup`),
    pero eso depende de como se clique el resultado. Se cubren los dos casos.
    """
    try:
        with page.expect_popup(timeout=_ms(min(cfg.timeout_operacion_seg, 10))) as info:
            locator.first.click()
        nueva = info.value
        nueva.wait_for_load_state("domcontentloaded")
        log.info("El informe se abrio en una pestana nueva.")
        return nueva
    except ErrorTiempoPlaywright:
        # Sin popup: el clic ya se ejecuto y navego en la misma pestana.
        log.debug("Sin pestana nueva; el informe navego en la misma pagina.")
        page.wait_for_load_state("domcontentloaded")
        return page


# ---------------------------------------------------------------------------
# Paso 3: exportar el visual
# ---------------------------------------------------------------------------


def _exportar_visual(
    page: Page,
    cfg: Config,
    advertencias: list[str],
    captura: Path | None = None,
) -> tuple[Download, int | None]:
    """Menu del visual -> Exportar datos -> Datos con diseno actual -> Exportar."""
    page.set_default_timeout(_ms(cfg.timeout_operacion_seg))

    _verificar_pagina(page, cfg, advertencias)
    # Se lee ANTES de exportar: la pagina ya esta cargada y el dialogo aun no
    # tapa el tablero.
    no_matriculado = _leer_tarjeta_no_matriculado(page, cfg, advertencias)
    visual = _localizar_visual_tabular(page, cfg, advertencias)
    visual.click()

    # El boton "..." se busca DENTRO del visual elegido. Buscarlo en la pagina
    # devolvia el del primer visual del DOM, que no es el que queremos.
    visual.hover()
    mas_opciones = visual.get_by_test_id(sel.TESTID_VISUAL_MAS_OPCIONES)
    if not _visible(mas_opciones, cfg.timeout_operacion_seg):
        raise ErrorExportacionPowerBI(
            "El visual de tabla no expuso su boton '...' de mas opciones "
            f"('{sel.TESTID_VISUAL_MAS_OPCIONES}')."
        )
    mas_opciones.first.click()

    _, exportar = _primero_visible(
        [
            ("test-id pbimenu-item.Exportar datos", page.get_by_test_id(sel.TESTID_MENU_EXPORTAR)),
            ("menuitem ~exportar datos", page.get_by_role("menuitem", name=sel.RX_MENU_EXPORTAR)),
        ],
        cfg.timeout_operacion_seg,
        "el item 'Exportar datos' del menu del visual",
    )
    exportar.first.click()

    _esperar_dialogo(page, cfg)
    _marcar_diseno_actual(page, cfg, advertencias)

    if captura is not None:
        _capturar_dialogo(page, captura, advertencias)

    _, boton_exportar = _primero_visible(
        [
            ("test-id export-btn", page.get_by_test_id(sel.TESTID_BOTON_EXPORTAR)),
            ("boton ~exportar", page.get_by_role("button", name=sel.RX_MENU_EXPORTAR)),
        ],
        cfg.timeout_operacion_seg,
        "el boton 'Exportar' del dialogo",
    )

    log.info("Confirmando la exportacion; la descarga puede tardar.")
    try:
        with page.expect_download(timeout=_ms(cfg.timeout_descarga_seg)) as info:
            boton_exportar.first.click()
    except ErrorTiempoPlaywright as exc:
        raise ErrorExportacionPowerBI(
            f"Power BI no entrego la descarga en {cfg.timeout_descarga_seg}s. "
            "Con muchas filas puede hacer falta subir TIMEOUT_DESCARGA_SEG."
        ) from exc
    return info.value, no_matriculado


def leer_valor_de_tarjeta(etiqueta: str) -> int | None:
    """Extrae el numero de un aria-label de tarjeta. None si no lo hay.

    Formato real observado: 'NO MATRICULADO 1323.' -- nombre y valor en la misma
    etiqueta, con un punto final que es puntuacion. Una tarjeta sin dato dice
    'MATRICULADO (En blanco).'.

    Los separadores de miles se eliminan: en es-CO el separador es el punto, de
    modo que '1.323' y '1323' son el mismo numero, y el punto final tampoco
    estorba al quitarlos todos.
    """
    texto = (etiqueta or "").strip()
    if not texto or sel.RX_TARJETA_EN_BLANCO.search(texto):
        return None

    m = sel.RX_VALOR_TARJETA.search(texto)
    if not m:
        return None

    digitos = re.sub(r"[^\d]", "", m.group(1))
    return int(digitos) if digitos else None


def _etiqueta_de_tarjeta(tarjeta: Locator) -> str:
    """Texto de una tarjeta, mire donde mire Power BI.

    Comprobado el 24/08/2026: el contenedor con
    `aria-roledescription="Tarjeta"` NO lleva aria-label; la etiqueta cuelga de
    un `<svg>` descendiente:

        contenedor  aria-label = null
          svg       aria-label = "NO MATRICULADO 131."
          innerText = "131\\nNO MATRICULADO"

    Se prueban las tres fuentes en orden de fiabilidad. El innerText invierte el
    orden (valor primero), pero el parseo no depende del orden.
    """
    try:
        etiqueta = tarjeta.get_attribute("aria-label")
        if etiqueta and etiqueta.strip():
            return etiqueta.strip()

        interno = tarjeta.locator("[aria-label]")
        for i in range(min(interno.count(), 5)):
            valor = interno.nth(i).get_attribute("aria-label")
            if valor and valor.strip():
                return valor.strip()

        return (tarjeta.inner_text() or "").strip().replace("\n", " ")
    except ErrorPlaywright as exc:
        log.debug("No se pudo leer la etiqueta de una tarjeta: %s", exc)
        return ""


def _leer_tarjeta_no_matriculado(
    page: Page, cfg: Config, advertencias: list[str]
) -> int | None:
    """Valor de la tarjeta 'NO MATRICULADO' del tablero.

    Es el numero que va en el nombre del reporte (`REPORTE <fecha> #<valor>`).
    No aborta si no se encuentra: el .xlsx sigue siendo valido y el valor se
    puede deducir de sus filas. Deja advertencia para que se note.
    """
    tarjetas = page.locator(sel.css_tarjetas())
    if not _visible(tarjetas, cfg.timeout_operacion_seg):
        advertencias.append(
            "No se encontro ninguna tarjeta en el tablero; no se pudo leer "
            "'NO MATRICULADO'."
        )
        log.warning(advertencias[-1])
        return None

    etiquetas = [
        _etiqueta_de_tarjeta(tarjetas.nth(i)) for i in range(tarjetas.count())
    ]
    log.info("Tarjetas en la pagina: %s", etiquetas)

    for etiqueta in etiquetas:
        if not sel.RX_TARJETA_NO_MATRICULADO.search(etiqueta):
            continue
        valor = leer_valor_de_tarjeta(etiqueta)
        if valor is None:
            advertencias.append(
                f"La tarjeta 'NO MATRICULADO' se encontro pero sin valor legible: "
                f"{etiqueta!r}."
            )
            log.warning(advertencias[-1])
            return None
        log.info("Tarjeta 'NO MATRICULADO' = %d (de %r)", valor, etiqueta)
        return valor

    advertencias.append(
        f"Ninguna tarjeta se llama 'NO MATRICULADO'. Encontradas: {etiquetas}."
    )
    log.warning(advertencias[-1])
    return None


def _verificar_pagina(page: Page, cfg: Config, advertencias: list[str]) -> None:
    """Comprueba que el lienzo cargado es la pagina esperada.

    El lienzo se expone como
    `<exploration data-automation-type="exploration" aria-label="<pagina>">`.
    Un aviso aqui explica de golpe cualquier fallo posterior de localizacion.
    """
    lienzo = page.locator(f'[{sel.ATRIBUTO_AUTOMATIZACION}="{sel.TESTID_LIENZO}"]')
    if not _visible(lienzo, cfg.timeout_render_seg):
        advertencias.append("No se pudo identificar el lienzo del informe.")
        return

    etiqueta = (lienzo.first.get_attribute("aria-label") or "").strip()
    if etiqueta.casefold() == cfg.powerbi_pagina.strip().casefold():
        log.info("Pagina del informe confirmada: '%s'.", etiqueta)
    else:
        advertencias.append(
            f"El informe abrio en la pagina '{etiqueta}', no en "
            f"'{cfg.powerbi_pagina}' (POWERBI_PAGINA)."
        )
        log.warning(advertencias[-1])


def _localizar_visual_tabular(page: Page, cfg: Config, advertencias: list[str]) -> Locator:
    """Devuelve el visual de tabla/matriz del que hay que exportar.

    Se busca por `aria-roledescription`, no por titulo: es la propiedad que
    determina si "Datos con diseno actual" estara habilitado, y no depende de
    como se llame el visual. `POWERBI_VISUAL` solo desempata cuando la pagina
    trae mas de una tabla.

    Raises:
        ErrorExportacionPowerBI: si la pagina no tiene ningun visual tabular.
    """
    candidatos = page.locator(sel.css_visual_tabular())
    if not _visible(candidatos, cfg.timeout_render_seg):
        # Sin tablas no hay exportacion posible con diseno actual: listar lo que
        # si hay ahorra el viaje a la traza.
        todos = page.locator(sel.CSS_CONTENEDOR_VISUAL)
        etiquetas = _etiquetas_de(todos)
        raise ErrorExportacionPowerBI(
            "La pagina del informe no tiene ningun visual de tabla o matriz, "
            "que son los unicos que permiten exportar 'Datos con diseno "
            f"actual'. Visuales encontrados: {etiquetas or 'ninguno'}. "
            "Revisar POWERBI_URL_INFORME y POWERBI_PAGINA."
        )

    etiquetas = _etiquetas_de(candidatos)
    log.info("Visuales tabulares en la pagina: %s", etiquetas)

    buscado = cfg.powerbi_visual.strip().casefold()
    if buscado:
        for indice, etiqueta in enumerate(etiquetas):
            if buscado in etiqueta.strip().casefold():
                log.info("Visual elegido por POWERBI_VISUAL: '%s'.", etiqueta)
                return candidatos.nth(indice)

        advertencias.append(
            f"Ningun visual tabular coincide con POWERBI_VISUAL='{cfg.powerbi_visual}'; "
            f"se usa el primero ('{etiquetas[0]}'). Etiquetas disponibles: {etiquetas}."
        )
        log.warning(advertencias[-1])
    elif len(etiquetas) > 1:
        advertencias.append(
            f"La pagina tiene {len(etiquetas)} visuales tabulares {etiquetas} y "
            "POWERBI_VISUAL esta vacio; se usa el primero."
        )
        log.warning(advertencias[-1])

    return candidatos.first


def _etiquetas_de(locator: Locator) -> list[str]:
    """aria-labels de todos los elementos que casa un locator.

    Los titulos de visual llegan con espacios sobrantes (el de la tabla es
    literalmente 'REPORTE '), asi que se normalizan.
    """
    etiquetas: list[str] = []
    for i in range(locator.count()):
        etiquetas.append((locator.nth(i).get_attribute("aria-label") or "").strip())
    return etiquetas


def _capturar_dialogo(page: Page, destino: Path, advertencias: list[str]) -> None:
    """Guarda una imagen del dialogo de exportacion antes de confirmar.

    Sirve para auditar el paso que no quedo en la grabacion: si la opcion
    correcta estaba marcada o no se ve en la imagen, sin tener que estar
    mirando la pantalla en el instante exacto.
    """
    destino.parent.mkdir(parents=True, exist_ok=True)
    dialogo = page.get_by_role("dialog")
    try:
        if dialogo.count():
            dialogo.first.screenshot(path=str(destino))
        else:
            page.screenshot(path=str(destino))
        log.info("Captura del dialogo: %s", destino)
    except ErrorTiempoPlaywright as exc:
        # Una captura fallida no justifica perder la exportacion.
        advertencias.append(f"No se pudo capturar el dialogo: {exc}")


def _esperar_dialogo(page: Page, cfg: Config) -> None:
    """Espera a que el dialogo de exportacion este montado."""
    acciones = page.get_by_test_id(sel.TESTID_DIALOGO_ACCIONES)
    if _visible(acciones, cfg.timeout_operacion_seg):
        return
    if _visible(page.get_by_role("dialog"), 5):
        log.warning(
            "El dialogo de exportacion no expuso '%s'; se continua por rol.",
            sel.TESTID_DIALOGO_ACCIONES,
        )
        return
    raise ErrorExportacionPowerBI(
        "No se abrio el dialogo de exportacion tras pulsar 'Exportar datos'."
    )


def _marcar_diseno_actual(page: Page, cfg: Config, advertencias: list[str]) -> None:
    """Selecciona "Datos con diseno actual" en el dialogo de exportacion.

    Este paso no quedo en la grabacion y en la corrida del 21/08/2026 11:11 se
    salto en silencio: la exportacion salio con el defecto ("Datos resumidos"),
    entrego un archivo de una columna y aun asi termino con codigo 0. Por eso
    ahora cualquier duda aqui ABORTA: exportar otra cosa es peor que no
    exportar, porque el archivo equivocado sigue camino a la Fase 1.

    El input nativo del radio esta oculto (tamano cero), asi que no se puede
    esperar a que sea visible ni clicarlo: se actua sobre su `<label for=...>` y
    se comprueba el resultado leyendo el input.

    Raises:
        ErrorExportacionPowerBI: si la opcion no existe, esta deshabilitada o no
            se logra dejar marcada.
    """
    grupo = page.get_by_test_id(sel.TESTID_GRUPO_RADIOS)
    dialogo = page.get_by_role("dialog")
    ambito = grupo if grupo.count() else (dialogo if dialogo.count() else page)

    opcion = _radio_por_etiqueta(ambito, sel.RX_OPCION_DISENO_ACTUAL)
    if opcion is None:
        disponibles = _etiquetas_de(ambito.locator(sel.CSS_RADIO_EXPORTACION))
        raise ErrorExportacionPowerBI(
            "El dialogo de exportacion no ofrece 'Datos con diseno actual'. "
            f"Opciones encontradas: {disponibles or 'ninguna'}."
        )

    if opcion.is_disabled():
        raise ErrorExportacionPowerBI(
            "'Datos con diseno actual' esta DESHABILITADO: Power BI solo la "
            "ofrece para visuales de tabla y matriz, asi que el visual "
            "seleccionado no es el correcto. Con el defecto ('Datos resumidos') "
            "se descarga otra cosa, de modo que se aborta. Revisar "
            f"POWERBI_VISUAL (ahora '{cfg.powerbi_visual}') y POWERBI_PAGINA."
        )

    if opcion.is_checked():
        log.info("'Datos con diseno actual' ya estaba seleccionado.")
        return

    _clicar_radio_oculto(ambito, opcion)

    if not opcion.is_checked():
        raise ErrorExportacionPowerBI(
            "Se pulso 'Datos con diseno actual' pero el radio no quedo marcado; "
            "no se continua para no exportar el formato equivocado."
        )
    log.info("Seleccionado 'Datos con diseno actual'.")


def _radio_por_etiqueta(ambito: Locator | Page, patron: re.Pattern[str]) -> Locator | None:
    """Radio del grupo cuyo aria-label casa el patron, o None.

    No usa `_visible`: estos inputs estan ocultos por diseno, de modo que
    esperar visibilidad los descartaria todos -- que es exactamente lo que
    hacia fallar la deteccion.
    """
    radios = ambito.locator(sel.CSS_RADIO_EXPORTACION)
    for i in range(radios.count()):
        radio = radios.nth(i)
        if patron.search(radio.get_attribute("aria-label") or ""):
            return radio
    return None


def _clicar_radio_oculto(ambito: Locator | Page, radio: Locator) -> None:
    """Marca un radio cuyo input nativo no es clicable.

    Se clica la etiqueta asociada por `for`; si el input no tiene id, se recurre
    a `check(force=True)`, que salta las comprobaciones de accionabilidad.
    """
    id_input = radio.get_attribute("id")
    if id_input:
        etiqueta = ambito.locator(f'label[for="{id_input}"]')
        if etiqueta.count():
            etiqueta.first.click()
            return
    radio.check(force=True)


# ---------------------------------------------------------------------------
# Orquestacion
# ---------------------------------------------------------------------------


def verificar_estructura_descarga(archivo: Path) -> None:
    """Comprueba que el .xlsx descargado es de verdad el reporte esperado.

    La corrida del 21/08/2026 11:11 termino con codigo 0 y un archivo de 7,7 KiB
    con una sola columna: se habia exportado "Datos resumidos" de otro visual.
    Un fichero equivocado que pasa por bueno es peor que un fallo, porque
    alimenta la Fase 1 sin avisar. Esta comprobacion es el ultimo filtro.

    Raises:
        ErrorExportacionPowerBI: si las cabeceras no son las esperadas.
    """
    libro = load_workbook(archivo, read_only=True)
    try:
        nombre_hoja = libro.sheetnames[0]
        hoja = libro[nombre_hoja]
        primera = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
        cabeceras = tuple((c or "").strip() if isinstance(c, str) else c for c in primera)
    finally:
        libro.close()

    if cabeceras == COLUMNAS_ESPERADAS:
        log.info(
            "Estructura verificada: hoja '%s' con las %d columnas esperadas.",
            nombre_hoja,
            len(COLUMNAS_ESPERADAS),
        )
        return

    faltan = [c for c in COLUMNAS_ESPERADAS if c not in cabeceras]
    pista = ""
    if len(cabeceras) <= 1:
        # Sintoma exacto de "Datos resumidos": una sola columna y el pie arriba.
        pista = (
            " Una sola columna es la firma de una exportacion en 'Datos "
            "resumidos' o de un visual que no es la tabla del reporte."
        )
    raise ErrorExportacionPowerBI(
        f"El archivo descargado no tiene la estructura del reporte ({archivo}). "
        f"Cabeceras encontradas: {list(cabeceras)}. Faltan: {faltan}.{pista}"
    )


def contar_filas_datos(archivo: Path) -> int:
    """Filas de datos del .xlsx, sin cabecera ni pie 'Filtros aplicados:'."""
    from .constantes import MARCADOR_PIE_FILTROS

    libro = load_workbook(archivo, read_only=True)
    try:
        hoja = libro[libro.sheetnames[0]]
        total = 0
        for i, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            if i == 1:
                continue  # cabecera
            primera = fila[0] if fila else None
            if isinstance(primera, str) and primera.startswith(MARCADOR_PIE_FILTROS):
                continue  # pie de la exportacion, no es un registro
            if any(c is not None and str(c).strip() for c in fila):
                total += 1
        return total
    finally:
        libro.close()


def _contrastar_tarjeta_con_filas(
    no_matriculado: int | None, filas: int, advertencias: list[str]
) -> None:
    """Comprueba que la tarjeta y el .xlsx cuentan lo mismo.

    La vista de Power BI viene filtrada a VALIDACION != MATRICULADO, asi que la
    tarjeta 'NO MATRICULADO' y las filas de la tabla son el mismo conjunto.
    Verificado el 21/08/2026: tarjeta 1323, .xlsx 1323 filas. Confirmado por el
    dueno del proceso el 24/08/2026 con 131.

    Si difieren, el nombre del reporte en Drive quedaria mal, asi que se avisa.
    """
    if no_matriculado is None:
        advertencias.append(
            f"No se leyo la tarjeta 'NO MATRICULADO'. El .xlsx trae {filas} filas "
            "de datos, que es el valor que la etapa 2 usara como respaldo."
        )
        log.warning(advertencias[-1])
        return

    if no_matriculado == filas:
        log.info("Tarjeta y .xlsx coinciden: %d registros.", filas)
        return

    advertencias.append(
        f"La tarjeta 'NO MATRICULADO' dice {no_matriculado} pero el .xlsx trae "
        f"{filas} filas de datos. Deberian coincidir. Revisar si la exportacion "
        "quedo incompleta antes de usar ese numero en el nombre del reporte."
    )
    log.warning(advertencias[-1])


def ruta_destino_por_defecto(sello: datetime | None = None) -> Path:
    """`data/raw/reporte_moodle_sinu_<AAAAMMDD_HHMMSS>.xlsx`."""
    marca = (sello or datetime.now()).strftime("%Y%m%d_%H%M%S")
    return DIR_CRUDO / f"reporte_moodle_sinu_{marca}.xlsx"


def exportar_reporte(
    cfg: Config | None = None,
    *,
    destino: Path | None = None,
    headless: bool | None = None,
    con_traza: bool = False,
    con_captura: bool = False,
) -> ResultadoExportacion:
    """Descarga el reporte Moodle vs SINU desde Power BI.

    Args:
        cfg: configuracion efectiva; por defecto se lee de `config/.env`.
        destino: ruta del .xlsx a escribir; por defecto, `data/raw/` con sello.
        headless: fuerza el modo del navegador por encima de POWERBI_HEADLESS.
            Con MFA hay que ejecutar visible al menos una vez.
        con_traza: guarda una traza de Playwright en `logs/` para depurar
            selectores (`playwright show-trace <ruta>`).
        con_captura: guarda en `logs/` una imagen del dialogo de exportacion
            antes de confirmar, como evidencia de la opcion seleccionada.

    Raises:
        ErrorExportacionPowerBI: si falta configuracion, si un paso no se pudo
            localizar o si la descarga no llego.
    """
    cfg = cfg or Config.desde_entorno()

    # Sin perfil persistente el login es inevitable, asi que se exigen las
    # credenciales antes de arrancar Chromium. Con perfil, la sesion guardada
    # puede bastar: se deja que el flujo lo descubra.
    if not cfg.powerbi_perfil_navegador:
        verificar_credenciales(cfg)

    destino = Path(destino) if destino else ruta_destino_por_defecto()
    destino.parent.mkdir(parents=True, exist_ok=True)

    sin_cabeza = cfg.powerbi_headless if headless is None else headless
    advertencias: list[str] = []
    ruta_traza: Path | None = None
    ruta_captura: Path | None = None
    no_matriculado: int | None = None
    inicio = time.monotonic()

    # Un solo sello para traza y captura: asi los dos artefactos de una misma
    # corrida se emparejan de un vistazo en logs/.
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    if con_captura:
        ruta_captura = DIR_LOGS / f"dialogo_exportacion_{sello}.png"

    with sync_playwright() as pw:
        context, cerrar = _abrir_contexto(pw, cfg, sin_cabeza)
        if con_traza:
            DIR_LOGS.mkdir(parents=True, exist_ok=True)
            ruta_traza = DIR_LOGS / f"traza_powerbi_{sello}.zip"
            context.tracing.start(screenshots=True, snapshots=True, sources=True)

        try:
            context.set_default_timeout(_ms(cfg.timeout_operacion_seg))
            page = context.pages[0] if context.pages else context.new_page()

            _iniciar_sesion(page, cfg, advertencias)
            pagina_informe, ruta_navegacion = _abrir_informe(context, page, cfg)
            descarga, no_matriculado = _exportar_visual(
                pagina_informe, cfg, advertencias, ruta_captura
            )

            nombre_sugerido = descarga.suggested_filename
            descarga.save_as(destino)
            log.info("Descargado '%s' -> %s", nombre_sugerido, destino)
        except ErrorExportacionPowerBI:
            raise
        except ErrorTiempoPlaywright as exc:
            raise ErrorExportacionPowerBI(
                f"Timeout de Playwright durante la exportacion: {exc}"
            ) from exc
        except ErrorPlaywright as exc:
            # Cualquier otro fallo del navegador (navegacion abortada, contexto
            # destruido, pagina cerrada...). Sin esta rama la excepcion escapaba
            # cruda: el CLI solo atrapa ErrorExportacionPowerBI, asi que el
            # traceback iba a la consola y el log de la corrida quedaba sin
            # ninguna linea de error.
            raise ErrorExportacionPowerBI(
                f"Error de Playwright durante la exportacion: {exc}"
            ) from exc
        finally:
            if con_traza and ruta_traza is not None:
                context.tracing.stop(path=str(ruta_traza))
                log.info("Traza: playwright show-trace %s", ruta_traza)
            cerrar()

    if not destino.is_file() or destino.stat().st_size == 0:
        raise ErrorExportacionPowerBI(f"La descarga quedo vacia o no se escribio: {destino}")

    verificar_estructura_descarga(destino)
    filas_datos = contar_filas_datos(destino)
    _contrastar_tarjeta_con_filas(no_matriculado, filas_datos, advertencias)

    return ResultadoExportacion(
        archivo=destino,
        bytes=destino.stat().st_size,
        segundos=round(time.monotonic() - inicio, 1),
        nombre_sugerido=nombre_sugerido,
        ruta_navegacion=ruta_navegacion,
        advertencias=advertencias,
        traza=ruta_traza,
        captura_dialogo=ruta_captura if ruta_captura and ruta_captura.is_file() else None,
        no_matriculado=no_matriculado,
        filas_datos=filas_datos,
    )


#: El arranque del navegador vive en `navegador.py`: lo comparten esta etapa y
#: la 2, para que las lecciones sobre idioma, viewport y perfil persistente no
#: haya que repetirlas (ni olvidarlas) en cada modulo.
_abrir_contexto = abrir_contexto
