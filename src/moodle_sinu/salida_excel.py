"""Escritura del resultado de validacion en una copia local del .xlsx.

Sirve como verificacion visual de la Fase 1 sin depender todavia de las APIs
de Google. La escritura definitiva en Google Sheets (etapa 5) reutiliza el
mismo `ResultadoValidacion` y la misma paleta.

Esta copia es ademas **el archivo que se sube a Drive**, asi que aqui termina
la rutina de limpieza: marcado de calidad y, encima, el orden A-Z por
COD_PERIODO que agrupa los periodos en bloques continuos (ver `orden_periodo`).
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .constantes import COL_VALIDACION_RPA, Color
from .modelos import ResultadoValidacion
from .orden_periodo import ordenar_hoja_por_periodo

log = logging.getLogger(__name__)


def _relleno(color: Color) -> PatternFill:
    return PatternFill(start_color=color.value, end_color=color.value, fill_type="solid")


def escribir_copia_coloreada(
    resultado: ResultadoValidacion,
    destino: str | Path,
    *,
    hoja: str | None = None,
    ordenar_periodo: bool = True,
) -> Path:
    """Escribe una copia del reporte con los colores y la columna VALIDACION_RPA.

    Precedencia de color: primero el color de fila (duplicado amarillo en la
    Fase 1; caso 1/2/3 en la Fase 2), y encima los colores de celda de las
    reglas 1 y 2. Asi una celda azul sigue distinguiendose sobre su fila.

    Args:
        ordenar_periodo: ordena la tabla completa A-Z por COD_PERIODO como
            ultimo paso, para que los periodos salgan en bloques continuos.
            Se hace **despues** de pintar, de modo que cada celda conserva su
            color, y **antes** de guardar, de modo que lo que se sube a Sheets
            ya va ordenado. Como los numeros de fila cambian, `resultado` se
            renumera en el sitio: sus `fila` siguen apuntando al archivo que
            esta funcion acaba de escribir.
    """
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(resultado.archivo)
    try:
        ws = wb[hoja] if hoja else wb[resultado.hoja]

        # --- columna nueva VALIDACION_RPA ---
        col_rpa = ws.max_column + 1
        letra_rpa = get_column_letter(col_rpa)
        celda_encabezado = ws.cell(row=1, column=col_rpa, value=COL_VALIDACION_RPA)
        celda_encabezado.font = Font(bold=True)
        ws.column_dimensions[letra_rpa].width = 32

        ultima_col = col_rpa

        # --- 1) color de fila ---
        for fila in resultado.filas:
            if fila.color_fila is None:
                continue
            relleno = _relleno(fila.color_fila)
            for col in range(1, ultima_col + 1):
                ws.cell(row=fila.fila, column=col).fill = relleno

        # --- 2) color de celda (se superpone al de fila) ---
        for hallazgo in resultado.hallazgos:
            if hallazgo.columna == "?":
                continue
            ws[f"{hallazgo.columna}{hallazgo.fila}"].fill = _relleno(hallazgo.color)

        # --- 3) etiqueta de validacion / motivo de omision ---
        for fila in resultado.filas:
            if fila.validacion_rpa is not None:
                texto = fila.validacion_rpa.value
            elif fila.omitida:
                texto = f"OMITIDA: {fila.motivo_omision}"
            else:
                texto = ""  # pendiente de Fase 2
            ws.cell(row=fila.fila, column=col_rpa, value=texto)

        # --- 4) orden A-Z por COD_PERIODO, con los colores ya puestos ---
        if ordenar_periodo:
            ordenar_hoja_por_periodo(ws, resultado, ultima_col=ultima_col)

        wb.save(destino)
        log.info("Copia coloreada escrita en %s", destino)
        return destino
    finally:
        wb.close()
