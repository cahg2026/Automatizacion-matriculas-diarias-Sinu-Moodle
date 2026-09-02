#!/usr/bin/env python3
"""
Extrae la lista de cedulas (IDENTIFICACION) unicas, en orden de aparicion,
desde una hoja de un Excel de matricula CUN.

Uso:
    python3 extract_cedulas.py <ruta.xlsx> [--sheet "Hoja1"] [--column "IDENTIFICACION"]

Imprime una cedula por linea (texto plano, sin encabezado), lista para
recorrerlas una por una en el sistema academico.
"""
import argparse
import sys

try:
    import openpyxl
except ImportError:
    print("Falta la libreria openpyxl. Instala con: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(1)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", help="Ruta al archivo .xlsx de matricula")
    parser.add_argument("--sheet", default=None, help="Nombre de la hoja a usar (por defecto: la primera)")
    parser.add_argument("--column", default="IDENTIFICACION", help="Nombre de la columna con la cedula")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]

    header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        col_idx = header.index(args.column)
    except ValueError:
        print(f"No se encontro la columna '{args.column}' en la hoja '{ws.title}'. Columnas disponibles: {header}", file=sys.stderr)
        sys.exit(1)

    seen = set()
    ordered = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if col_idx >= len(row):
            continue
        val = row[col_idx]
        if val is None:
            continue
        cid = str(val).strip()
        if cid and cid not in seen:
            seen.add(cid)
            ordered.append(cid)

    for cid in ordered:
        print(cid)

    print(f"# total filas: {ws.max_row - 1} | cedulas unicas: {len(ordered)} | hoja: {ws.title}", file=sys.stderr)


if __name__ == "__main__":
    main()
